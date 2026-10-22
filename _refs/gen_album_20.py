# -*- coding: utf-8 -*-
# 智多星平台 20 任务图册（电子版）：每行一个任务＋截图占位框；真实总览图=封面地图
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RPA = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "RPA")
OUT = "智多星平台任务图册_电子版_v1.0_20261023.docx"
PROJECT = "智多星RPA基础命令"

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)

def T(txt, size=10.5, bold=False, color="333333", center=False, before=0, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = "微软雅黑"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return p

def box(hint, h=6.2):
    tb = doc.add_table(rows=1, cols=1); tb.style = "Table Grid"
    cell = tb.cell(0, 0); cell.width = Cm(16.6)
    tr = tb.rows[0]._tr; trPr = tr.get_or_add_trPr()
    el = OxmlElement("w:trHeight"); el.set(qn("w:val"), str(int(h*567))); el.set(qn("w:hRule"), "atLeast")
    trPr.append(el)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(hint); r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string("999999")
    r.font.name = "微软雅黑"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

def head(txt):
    tb = doc.add_table(rows=1, cols=1); tb.style = "Table Grid"
    cell = tb.cell(0, 0); cell.width = Cm(16.6)
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "BF360C")
    cell._tc.get_or_add_tcPr().append(sh)
    p = cell.paragraphs[0]
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string("FFFFFF")
    r.font.name = "微软雅黑"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# ── 封面 ──
doc.add_page_break()
T("智多星 RPA 基础命令 · 20 任务图册", 20, True, "1A5276", True, 6, 4)
T("电子版 · 对照平台逐个过：任务状态变「已完成」＝过；每过一张，把任务页截图贴进本行方框", 11, False, "555555", True, 0, 6)
T("四大板块对应纸质册《机器人实验记录册》四个暂停点：变量操作→⑬｜流程设计操作→⑭｜数据操作→⑮｜Office操作→⑯", 10.5, True, "C0390B", True, 0, 10)
pic = os.path.join(RPA, "img", "zhiduoxing_20task_overview.png")
if os.path.exists(pic):
    doc.add_picture(pic, width=Cm(16.6))
    T("▲ 平台总览（真实截图 2026-10-22）：四板块 20 任务——状态/实战次数就是完成凭证", 9.5, False, "777777", True, 2, 8)

SECTS = [
 ("变量操作 · 6 任务 → 纸质册暂停点⑬", [
  ("任务1·普通变量命令介绍", "文档", "看普通变量怎么建、怎么放数——对课上「储物格」"),
  ("任务2·普通变量知识测验", "作业", "独立做；错了回任务1文档再过一遍"),
  ("任务3·列表变量命令介绍", "文档", "一排格子——对课上「快递柜」；纸质册⑬问答①等你"),
  ("任务4·列表变量知识测验", "作业", "独立做"),
  ("任务5·字典变量命令介绍", "文档", "带标签的柜子组——「键→值」"),
  ("任务6·字典变量知识测验", "作业", "独立做；纸质册⑬问答②写一条你自己的键值对"),
 ]),
 ("流程设计操作 · 6 任务 → 纸质册暂停点⑭", [
  ("任务1·条件命令介绍", "文档", "「如果…就…」——对外卖图纸的岔路口那行"),
  ("任务2·条件命令知识测验", "作业", "独立做"),
  ("任务3·循环命令介绍", "文档", "次数＋肚子——对 A1–A7 写 1–7"),
  ("任务4·循环命令知识测验", "作业", "独立做"),
  ("任务5·文件命令介绍", "文档", "课上没讲——自学，这是你的家底"),
  ("任务6·文件命令知识测验", "作业", "独立做；纸质册⑭问答③写它最可能帮你干的一件事"),
 ]),
 ("数据操作 · 6 任务 → 纸质册暂停点⑮", [
  ("任务1·数值命令介绍", "文档", "数值的加减乘除——39.9＋10 算得出"),
  ("任务2·数值命令知识测验", "作业", "独立做"),
  ("任务3·字符串命令介绍", "文档", "文字的拼接——「2501」＋「班」"),
  ("任务4·字符串命令知识测验", "作业", "独立做"),
  ("任务5·日期命令介绍", "文档", "日期的取与算——自学"),
  ("任务6·日期命令知识测验", "作业", "独立做；纸质册⑮问答③写一个用得上日期的活"),
 ]),
 ("Office操作 · 2 任务 → 纸质册暂停点⑯（最简单的一关）", [
  ("任务1·Excel命令介绍", "文档", "新建/写入/读格——就是大案例记账表用的那几条"),
  ("任务2·Excel命令知识测验", "作业", "过了就去开工大案例：5 笔消费、机器人填金额、弹合计"),
 ]),
]
n = 0
for title, tasks in SECTS:
    doc.add_page_break()
    head(title); T("", 6)
    for name, typ, note in tasks:
        n += 1
        T("%s（%s）——%s" % (name, typ, note), 10.5, False, "333333", False, 4, 2)
        box("📷 截图贴这里：本任务页整屏（状态「已完成」＋实战次数可见）")
T("—— 全 20 任务完成后：四板块状态全绿，纸质册⑬–⑯问答写全，凭大案例放行条进机房 ——", 10.5, True, "1A5276", True, 10, 0)
doc.save(os.path.join(RPA, OUT))
print("生成", OUT, "| 20 任务", len(SECTS), "节")
