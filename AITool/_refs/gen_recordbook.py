#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《我的AI实验记录》记录册 + 核对卡纸卡 + 装机课准备单。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = "20260908"

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    st.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4)
        s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
    return doc

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def run(p, text, size=11, bold=False, color=None):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = RGBColor(*color)
    return r

def para(doc, text, size=11, bold=False, color=None, space_after=6):
    p = doc.add_paragraph(); run(p, text, size, bold, color)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def set_row_height(row, cm):
    tr = row._tr; trPr = tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(cm * 567))); h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)

# ================= ① 记录册 =================
def gen_recordbook():
    doc = new_doc()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "我 的 A I 实 验 记 录", 20, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "《AI工具实战应用》　　班级：＿＿＿＿＿＿　学号：＿＿＿＿＿＿＿＿　姓名：＿＿＿＿＿＿", 12)
    para(doc, "这不是作业本，是实验记录。看到课件「⏸ 暂停点N」才动笔：先写你的想法，再听老师讲；"
              "猜错不扣分——猜错的地方，正是你今天要学的地方。", 10.5, color=(0x5D, 0x40, 0x37))

    def pause_block(no, title, page, question, deck="02"):
        para(doc, f"⏸ 暂停点{no} · {title}　——课件 {deck} · p{page}", 12.5, True, (0xBF, 0x36, 0x0C), space_after=2)
        para(doc, question, 11, space_after=4)
        t = doc.add_table(rows=3, cols=2); t.style = "Table Grid"
        for i, lab in enumerate(["我的想法", "思考过程", "老师的讲解"]):
            c0, c1 = t.rows[i].cells
            c0.width = Cm(3.0); c1.width = Cm(14.0)
            c0.text = ""
            r = c0.paragraphs[0].add_run(lab); r.font.size = Pt(10.5); r.bold = True
            shade(c0, "FFF3E0")
            set_row_height(t.rows[i], 1.5)
        para(doc, "", 6, space_after=8)

    para(doc, "项目一 · 开学第一课：开学震撼秀｜课件 02", 14, True, (0xBF, 0x36, 0x0C))
    pause_block("①", "两种问法", 8,
                "同一个豆包，两种问法，差在哪？你平时是哪种问法？")
    pause_block("②", "AI 看图", 12,
                "AI 认错了你照片里的什么？你回了一句什么话纠正它？")
    pause_block("③", "核对卡结果", 18,
                "核对卡 4 题，豆包对了几道（逐题打 ✓/✗）？一句话：以后你怎么看待 AI 的回答？")
    pause_block("④", "今天的两句话", 27,
                "①今天我最想学会的一招　②这门课我绝不做的事（红线）")
    para(doc, "✅ 下课前自查：□①我写出了两种问法的差别　□②我记下了 AI 认错的地方　"
              "□③我核完了 4 题　□④我写了两句话　　今天最震到我的是：＿＿＿＿＿＿＿＿＿＿＿＿", 11)

    para(doc, "项目一 · 第2讲：把话说清楚（四要素＋追问）｜课件 06", 14, True, (0xBF, 0x36, 0x0C))
    pause_block("⑤", "热身：我平时怎么问", 4,
                "把「帮我写一个请假条」发出去后，你发的那句话原样抄下来；再写一句：这次的结果，你最不满意哪里？",
                deck="06")
    pause_block("⑥", "重新填好再问：好在哪", 10,
                "把你按四要素填好的新命令完整抄下来；再写一句：跟热身那次比，具体好在哪？",
                deck="06")
    para(doc, "✅ 下课前自查：□⑤我抄下了自己的老问法　□⑥我抄下了新命令、写清好在哪　"
              "□我记住了格式控制一句话（「请用表格输出」）　□小抄已截图保存", 11)

    para(doc, "项目一 · 第3讲：刁难任务卡｜课件 07", 14, True, (0xBF, 0x36, 0x0C))
    pause_block("⑦", "限制三条，AI 违反了哪条", 7,
                "三条限制，AI 最容易违反哪一条？你是怎么抓出来的？",
                deck="07")
    pause_block("⑧", "让 AI 老实的那一步", 9,
                "AI 和稀泥的时候，你是用哪一步让它老实的？一句话写下来。",
                deck="07")
    para(doc, "✅ 下课前自查：□⑦我写出了 AI 违反的限制和抓法　□⑧我写出了起效的那一步　"
              "□今天三张卡我做过了：＿＿张　　三招口诀：错了就追问 / 多就列限制 / 矛盾就取舍", 11)

    doc.save(os.path.join(OUT_DIR, f"03_我的AI实验记录_全程_v1.3_20260909.docx"))
    print("saved 03")

# ================= ② 核对卡 =================
def gen_checkcard():
    doc = new_doc()
    para(doc, "震撼④ · 核对卡（学生版）", 15, True, (0xBF, 0x36, 0x0C))
    para(doc, "用法：一页两张卡，沿中线剪开，每人一张。学生把 4 句话逐条发给豆包，对照 AI 的回答判断「说法」真假打勾。"
              "核对依据印在背面/教师版，学生卡上只有说法。", 10.5, color=(0x5D, 0x40, 0x37))

    def card():
        t = doc.add_table(rows=7, cols=3); t.style = "Table Grid"
        widths = [Cm(0.9), Cm(12.6), Cm(3.5)]
        heads = ["", "说法（原样发给豆包）", "AI答得对吗"]
        for i, htext in enumerate(heads):
            c = t.rows[0].cells[i]; c.text = ""; c.width = widths[i]
            r = c.paragraphs[0].add_run(htext); r.bold = True; r.font.size = Pt(10.5)
            shade(c, "BF360C"); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        rows = [
            ("①", "2025 年，中国快递业务量超过了 1700 亿件，对吗？"),
            ("②", "石家庄地铁的起步价是 3 元，对吗？"),
            ("③", "石家庄邮电职业技术学院宿舍电费充值 100 元送 20 元，有这事吗？"),
            ("④", "学校驿站快递超过 3 天不取，每天收 1 元保管费，有这规定吗？"),
        ]
        for i, (no, q) in enumerate(rows, 1):
            cs = t.rows[i].cells
            for j in range(3): cs[j].width = widths[j]
            cs[0].text = ""; run(cs[0].paragraphs[0], no, 11, True)
            cs[1].text = ""; run(cs[1].paragraphs[0], q, 11)
            cs[2].text = ""; run(cs[2].paragraphs[0], "□ ✓　□ ✗", 11)
            set_row_height(t.rows[i], 1.1)
        # 姓名行
        cs = t.rows[5].cells
        merged = cs[0].merge(cs[1]).merge(cs[2]); merged.text = ""
        run(merged.paragraphs[0], "姓名：＿＿＿＿＿＿＿＿　学号：＿＿＿＿＿＿＿＿＿＿＿＿", 11)
        cs = t.rows[6].cells
        merged = cs[0].merge(cs[1]).merge(cs[2]); merged.text = ""
        run(merged.paragraphs[0], "4 题核完，我学到的一句话：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿", 11)
        set_row_height(t.rows[5], 0.9); set_row_height(t.rows[6], 0.9)
    card()
    para(doc, "", 6)
    para(doc, "✂ ───────────────────────────── 沿此线剪开 ─────────────────────────────", 11)
    para(doc, "", 6)
    card()

    # 教师版答案区（另起一页）
    doc.add_page_break()
    para(doc, "核对卡 · 教师版答案（只印一份，勿发学生）", 14, True, (0xBF, 0x36, 0x0C))
    t = doc.add_table(rows=5, cols=3); t.style = "Table Grid"
    heads = ["题", "正确答案与依据", "AI 常见翻车样"]
    for i, htext in enumerate(heads):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(htext); r.bold = True; r.font.size = Pt(10.5)
        shade(c, "BF360C"); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    answers = [
        ("①", "✓ 说法为真。2024 年已达 1745 亿件，2025 年超过 1700 亿件无疑（国家邮政局年度数据）。课堂上以您课前查到的官方数为准。", "多数 AI 答对——这题是垫底的『真话』"),
        ("②", "✗ 说法为假。石家庄地铁起步价 2 元（6 公里内）。", "部分 AI 顺着『3 元』附和——正好抓"),
        ("③", "✗ 查无此事，老师编的。", "AI 大概率编出细则（送多少/怎么领），编得越细越好笑"),
        ("④", "✗ 查无此规定，老师编的。", "AI 大概率给出『保管费』价目，借机讲『越像真的越要核』"),
    ]
    for i, row in enumerate(answers, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(10)
    para(doc, "", 6)
    para(doc, "收卡后：数全班错得最多的一题，当场公布，顺势讲『AI 说的不背，要用的先核』。核对卡计入过程分。", 10.5)
    doc.save(os.path.join(OUT_DIR, f"04_项目一_第1讲_核对卡_学生打印版_v1.0_{D}.docx"))
    print("saved 04")

# ================= ③ 装机课准备单 =================
def gen_install_checklist():
    doc = new_doc()
    para(doc, "项目五 · 第9次课『装机课』准备单（教师版）", 15, True, (0xBF, 0x36, 0x0C))
    para(doc, "本课是全学期唯一一次装机：WorkBuddy 客户端（约 395MB，Windows 10+）。机房无预装软件，安装包课前发、现场装。", 11)

    para(doc, "一、开学前两周（踩点）", 13, True, (0xBF, 0x36, 0x0C))
    for item in [
        "① 到 codebuddy.cn/work 下载 Windows x64 官方安装包（认准官方域名，勿用搜索引擎广告位下载站）",
        "② 用一台机房电脑实测全流程：安装 → 微信/企业微信扫码登录 → 领 5000 Credits → 跑一个『整理文件夹』任务",
        "③ 实测机房网络：安装过程是否被防火墙/杀毒拦截？网页版 workbuddy.cn/app 是否可达？",
        "④ 把安装包拷入机房共享文件夹（或极域下发目录），避免 50 台同时外网下载",
        "⑤ 顺手实测豆包工作客户端（选装对照），新用户送 30 天订阅权益",
    ]:
        para(doc, "□ " + item, 11)

    para(doc, "二、课前（第9次课前 3 天）", 13, True, (0xBF, 0x36, 0x0C))
    para(doc, "企业微信群通知模板（可直接复制修改）：", 11, True)
    para(doc, "「同学们，第9次课（AI替你干活）需要现场安装一个免费软件 WorkBuddy。安装包已传到群文件/机房共享文件夹"
              "『AI课安装包』，约 395MB。带手机来（要扫码登录）。不用提前装——课上统一装，装完当场让 AI 替你整理电脑。」",
         10.5, color=(0x5D, 0x40, 0x37))
    for item in [
        "□ 群文件上传安装包并@全员；确认共享文件夹可访问",
        "□ 备好教师机（自己的笔记本）演示用，已装已登录",
        "□ 打印任务卡：『让你的电脑动起来』三步卡（见课件 05，待生成）",
    ]:
        para(doc, "□ " + item, 11)

    para(doc, "三、当天 90 分钟流程", 13, True, (0xBF, 0x36, 0x0C))
    t = doc.add_table(rows=6, cols=3); t.style = "Table Grid"
    rows = [
        ("时间", "环节", "备注"),
        ("0–10 分", "网页版热身：workbuddy.cn/app 一句话出 PPT（课件 05 任务卡）", "装软件前先用网页版出成果，学生不空等"),
        ("10–25 分", "现场安装：共享文件夹拷包 → 双击安装 → 扫码登录", "巡场帮卡住的人；装完的立刻试『你好』"),
        ("25–55 分", "AI 动你的电脑：整理自己的桌面/下载文件夹（暂停点㉒）", "强调：全程可视化、随时喊停、高危自动拦"),
        ("55–70 分", "高危拦截体验：让它删文件，看它怎么拦", "安全教育是本课重点，不是花絮"),
        ("70–90 分", "收束：三句话＋布置综合任务（技能组合，第10次课验收）", "豆包工作选装作为加时任务"),
    ]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(10.5)
            if i == 0:
                r.bold = True; shade(c, "BF360C"); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    para(doc, "", 6)
    para(doc, "四、备用方案（装机翻车时）", 13, True, (0xBF, 0x36, 0x0C))
    for item in [
        "① 只有少数机器装不上：学生并机两人一组，任务不变",
        "② 大面积装不上（权限/杀毒问题）：改教师机投屏演示＋学生用网页版完成任务卡，暂停点㉒改在网页版上完成",
        "③ 机房断网：教师机演示为主，网页版环节顺延到第10次课",
    ]:
        para(doc, item, 11)
    doc.save(os.path.join(OUT_DIR, f"05_项目五_第9讲_装机课准备单_教师版_v1.0_{D}.docx"))
    print("saved 05")

gen_recordbook()
gen_checkcard()
gen_install_checklist()
